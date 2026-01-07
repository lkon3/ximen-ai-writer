import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, simpledialog, messagebox
import os
import json
import re
import shutil
import docx
import openai
import threading
import logging
from logging.handlers import RotatingFileHandler
import datetime
import random
import requests
from typing import Dict, List, Any, Callable
import hashlib
import base64

# ================================== AI WRITER V10.1 (UI/UX Redesign) ==================================
# 服务器后端API的地址
SERVER_URL = "http://47.99.123.64:5000"
# << 新增: 定义设定文件的专用名称 >>
SETTINGS_FILENAME = "__settings__.txt"


# ================================== 对话框类 (Minor theme updates will apply automatically) ==================================
class ForgotPasswordWindow(ctk.CTkToplevel):
    def __init__(self, parent, x, y):
        super().__init__(parent)
        self.parent = parent
        self.title("重置密码")
        self.geometry(f"450x680+{x}+{y}")
        self.transient(parent)
        self.grab_set()
        
        self.theme = parent.theme

        self.configure(fg_color=self.theme['bg'])
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.main_frame = ctk.CTkFrame(self, fg_color=self.theme['card'], corner_radius=16)
        self.main_frame.grid(row=0, column=0, padx=30, pady=30, sticky="nsew")
        self.main_frame.grid_columnconfigure(0, weight=1)

        back_button = ctk.CTkButton(self.main_frame, text="← 返回登录", command=self.go_back, font=("Microsoft YaHei UI", 13), fg_color="transparent", text_color=self.theme['text_light'], hover=False)
        back_button.grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")

        ctk.CTkLabel(self.main_frame, text="重置您的密码", font=("Microsoft YaHei UI", 26, "bold"), text_color=self.theme['text_h1']).grid(row=1, column=0, pady=(10, 30))

        ctk.CTkLabel(self.main_frame, text="邮箱地址", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=2, column=0, padx=30, sticky="ew")
        self.email_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.email_entry.grid(row=3, column=0, padx=30, pady=(5, 15), sticky="ew")

        ctk.CTkLabel(self.main_frame, text="验证码", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=4, column=0, padx=30, sticky="ew")
        self.code_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        self.code_frame.grid(row=5, column=0, padx=30, pady=5, sticky="ew")
        self.code_frame.grid_columnconfigure(0, weight=1)
        self.verification_code_entry = ctk.CTkEntry(self.code_frame, font=("Microsoft YaHei UI", 14), corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.verification_code_entry.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        self.send_code_button = ctk.CTkButton(self.code_frame, text="发送", command=self.handle_send_verification_code, font=("Microsoft YaHei UI", 13, "bold"), width=80, height=40, corner_radius=8, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'])
        self.send_code_button.grid(row=0, column=1, sticky="e")
        self.send_code_countdown = 0

        ctk.CTkLabel(self.main_frame, text="设置新密码", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=6, column=0, padx=30, sticky="ew")
        self.password_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), show="*", corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.password_entry.grid(row=7, column=0, padx=30, pady=(5, 15), sticky="ew")

        self.status_label = ctk.CTkLabel(self.main_frame, text="", font=("Microsoft YaHei UI", 13), text_color=self.theme['danger'], wraplength=300)
        self.status_label.grid(row=8, column=0, padx=30, pady=(15, 10), sticky="ew")

        self.reset_button = ctk.CTkButton(self.main_frame, text="确认重置密码", command=self.handle_reset_password, font=("Microsoft YaHei UI", 14, "bold"), height=45, corner_radius=8, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'])
        self.reset_button.grid(row=9, column=0, padx=30, pady=(10, 40), sticky="ew")

    def go_back(self):
        x = self.winfo_x(); y = self.winfo_y()
        self.parent.show_self(x, y); self.destroy()
    def set_status(self, message, is_error=True):
        text_color = self.theme['danger'] if is_error else self.theme['primary']
        self.status_label.configure(text=message, text_color=text_color)
    def toggle_buttons(self, enabled=True):
        state = "normal" if enabled else "disabled"
        self.reset_button.configure(state=state)
        if enabled and self.send_code_countdown <= 0: self.send_code_button.configure(state="normal")
        else: self.send_code_button.configure(state="disabled")
    def is_valid_email_format_client(self, email):
        return re.match(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$", email) is not None
    def handle_send_verification_code(self):
        email = self.email_entry.get().lower()
        if not email: self.set_status("请输入您注册时使用的邮箱"); return
        if not self.is_valid_email_format_client(email): self.set_status("邮箱格式不正确"); return
        self.set_status("正在发送验证码...", is_error=False)
        self.send_code_button.configure(state="disabled", text="发送中...")
        threading.Thread(target=self._do_send_verification_code, args=(email,), daemon=True).start()
    def _do_send_verification_code(self, email):
        try:
            response = requests.post(f"{SERVER_URL}/send_verification_code", json={"email": email}, timeout=15)
            try: data = response.json(); message = data.get("message", "未知错误")
            except ValueError: message = f"服务器错误 (代码: {response.status_code})"
            if response.status_code == 200 and data.get("success"):
                self.after(0, self.set_status, message, False)
                self.after(0, self._start_countdown, 60)
            else:
                self.after(0, self.set_status, message, True)
                self.after(0, lambda: self.send_code_button.configure(state="normal", text="发送"))
        except requests.RequestException:
            self.after(0, self.set_status, "无法连接到服务器或网络超时", True)
            self.after(0, lambda: self.send_code_button.configure(state="normal", text="发送"))
    def _start_countdown(self, seconds):
        self.send_code_countdown = seconds; self._update_countdown_button()
    def _update_countdown_button(self):
        if self.send_code_countdown > 0:
            self.send_code_button.configure(state="disabled", text=f"{self.send_code_countdown}s")
            self.send_code_countdown -= 1
            self.after(1000, self._update_countdown_button)
        else: self.send_code_button.configure(state="normal", text="发送")
    def handle_reset_password(self):
        email = self.email_entry.get().lower(); code = self.verification_code_entry.get(); new_password = self.password_entry.get()
        if not email or not code or not new_password: self.set_status("所有字段均为必填项"); return
        if not self.is_valid_email_format_client(email): self.set_status("邮箱格式不正确"); return
        self.set_status("正在重置密码...", is_error=False); self.toggle_buttons(False)
        threading.Thread(target=self._do_reset_password, args=(email, code, new_password), daemon=True).start()
    def _do_reset_password(self, email, code, new_password):
        try:
            payload = {"email": email, "code": code, "new_password": new_password}
            response = requests.post(f"{SERVER_URL}/reset_password", json=payload, timeout=10)
            data = response.json(); is_success = data.get("success", False)
            self.after(0, self.set_status, data.get("message", "重置失败"), not is_success)
            if is_success: self.after(2000, lambda: self.set_status("重置成功！请返回登录。", is_error=False))
        except requests.RequestException: self.after(0, self.set_status, "无法连接到服务器")
        finally: self.after(0, self.toggle_buttons, True)

class RegisterWindow(ctk.CTkToplevel):
    def __init__(self, parent, x, y):
        super().__init__(parent)
        self.parent = parent
        self.title("注册新账户")
        self.geometry(f"450x680+{x}+{y}")
        self.transient(parent)
        self.grab_set()
        
        self.theme = parent.theme

        self.configure(fg_color=self.theme['bg'])
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.main_frame = ctk.CTkFrame(self, fg_color=self.theme['card'], corner_radius=16)
        self.main_frame.grid(row=0, column=0, padx=30, pady=30, sticky="nsew")
        self.main_frame.grid_columnconfigure(0, weight=1)

        back_button = ctk.CTkButton(self.main_frame, text="← 返回登录", command=self.go_back, font=("Microsoft YaHei UI", 13), fg_color="transparent", text_color=self.theme['text_light'], hover=False)
        back_button.grid(row=0, column=0, padx=20, pady=(20, 0), sticky="w")

        ctk.CTkLabel(self.main_frame, text="创建您的账户", font=("Microsoft YaHei UI", 26, "bold"), text_color=self.theme['text_h1']).grid(row=1, column=0, pady=(10, 30))

        ctk.CTkLabel(self.main_frame, text="邮箱地址", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=2, column=0, padx=30, sticky="ew")
        self.email_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.email_entry.grid(row=3, column=0, padx=30, pady=(5, 15), sticky="ew")

        ctk.CTkLabel(self.main_frame, text="设置密码", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=4, column=0, padx=30, sticky="ew")
        self.password_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), show="*", corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.password_entry.grid(row=5, column=0, padx=30, pady=(5, 15), sticky="ew")

        ctk.CTkLabel(self.main_frame, text="验证码", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=6, column=0, padx=30, sticky="ew")
        self.code_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        self.code_frame.grid(row=7, column=0, padx=30, pady=5, sticky="ew")
        self.code_frame.grid_columnconfigure(0, weight=1)
        self.verification_code_entry = ctk.CTkEntry(self.code_frame, font=("Microsoft YaHei UI", 14), corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.verification_code_entry.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        self.send_code_button = ctk.CTkButton(self.code_frame, text="发送", command=self.handle_send_verification_code, font=("Microsoft YaHei UI", 13, "bold"), width=80, height=40, corner_radius=8, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'])
        self.send_code_button.grid(row=0, column=1, sticky="e")
        self.send_code_countdown = 0

        self.status_label = ctk.CTkLabel(self.main_frame, text="", font=("Microsoft YaHei UI", 13), text_color=self.theme['danger'], wraplength=300)
        self.status_label.grid(row=8, column=0, padx=30, pady=(15, 10), sticky="ew")

        self.register_button = ctk.CTkButton(self.main_frame, text="立即注册", command=self.handle_register, font=("Microsoft YaHei UI", 14, "bold"), height=45, corner_radius=8, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'])
        self.register_button.grid(row=9, column=0, padx=30, pady=(10, 40), sticky="ew")

    def go_back(self):
        x = self.winfo_x(); y = self.winfo_y()
        self.parent.show_self(x, y); self.destroy()
    def set_status(self, message, is_error=True):
        text_color = self.theme['danger'] if is_error else self.theme['primary']
        self.status_label.configure(text=message, text_color=text_color)
    def toggle_buttons(self, enabled=True):
        state = "normal" if enabled else "disabled"
        self.register_button.configure(state=state)
        if enabled and self.send_code_countdown <= 0: self.send_code_button.configure(state="normal")
        else: self.send_code_button.configure(state="disabled")
    def is_valid_email_format_client(self, email): return re.match(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$", email) is not None
    def handle_send_verification_code(self):
        email = self.email_entry.get().lower()
        if not email: self.set_status("邮箱不能为空"); return
        if not self.is_valid_email_format_client(email): self.set_status("邮箱格式不正确"); return
        self.set_status("正在发送验证码...", is_error=False); self.send_code_button.configure(state="disabled", text="发送中...")
        threading.Thread(target=self._do_send_verification_code, args=(email,), daemon=True).start()
    def _do_send_verification_code(self, email):
        try:
            response = requests.post(f"{SERVER_URL}/send_verification_code", json={"email": email}, timeout=15)
            try: data = response.json(); message = data.get("message", "未知错误")
            except ValueError: message = f"服务器错误 (代码: {response.status_code})"
            if response.status_code == 200 and data.get("success"):
                self.after(0, self.set_status, message, False); self.after(0, self._start_countdown, 60)
            else:
                self.after(0, self.set_status, message, True); self.after(0, lambda: self.send_code_button.configure(state="normal", text="发送"))
        except requests.RequestException:
            self.after(0, self.set_status, "无法连接到服务器或网络超时", True); self.after(0, lambda: self.send_code_button.configure(state="normal", text="发送"))
    def _start_countdown(self, seconds): self.send_code_countdown = seconds; self._update_countdown_button()
    def _update_countdown_button(self):
        if self.send_code_countdown > 0:
            self.send_code_button.configure(state="disabled", text=f"{self.send_code_countdown}s")
            self.send_code_countdown -= 1; self.after(1000, self._update_countdown_button)
        else: self.send_code_button.configure(state="normal", text="发送")
    def handle_register(self):
        email = self.email_entry.get().lower(); password = self.password_entry.get(); verification_code = self.verification_code_entry.get()
        if not email or not password or not verification_code: self.set_status("邮箱、密码和验证码不能为空"); return
        if not self.is_valid_email_format_client(email): self.set_status("邮箱格式不正确"); return
        self.set_status("注册中...", is_error=False); self.toggle_buttons(False)
        threading.Thread(target=self._do_register, args=(email, password, verification_code), daemon=True).start()
    def _do_register(self, email, password, verification_code):
        try:
            payload = {"email": email, "password": password, "code": verification_code}
            response = requests.post(f"{SERVER_URL}/register", json=payload, timeout=10)
            data = response.json(); is_success = data.get("success", False)
            self.after(0, self.set_status, data.get("message", "注册失败"), not is_success)
            if is_success: self.after(2000, lambda: self.set_status("注册成功！请返回登录。", is_error=False))
        except requests.RequestException: self.after(0, self.set_status, "无法连接到服务器")
        finally: self.after(0, self.toggle_buttons, True)

class LoginWindow(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.user_data = None
        self.register_win = None
        self.forgot_password_win = None

        self.title("西门写作")
        self.geometry("450x620")
        self.transient(parent)
        self.grab_set()
        self.protocol("WM_DELETE_WINDOW", self.on_closing)

        self.theme = parent.theme

        self.configure(fg_color=self.theme['bg'])
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.main_frame = ctk.CTkFrame(self, fg_color=self.theme['card'], corner_radius=16)
        self.main_frame.grid(row=0, column=0, padx=40, pady=40, sticky="nsew")
        self.main_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(self.main_frame, text="欢迎登录", font=("Microsoft YaHei UI", 28, "bold"), text_color=self.theme['text_h1']).grid(row=0, column=0, pady=(40, 10))
        ctk.CTkLabel(self.main_frame, text="使用您的邮箱和密码继续", font=("Microsoft YaHei UI", 14), text_color=self.theme['text_light']).grid(row=1, column=0, pady=(0, 30))

        ctk.CTkLabel(self.main_frame, text="邮箱地址", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=2, column=0, padx=30, sticky="ew")
        self.email_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), placeholder_text="user@example.com", corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.email_entry.grid(row=3, column=0, padx=30, pady=(5, 15), sticky="ew")

        ctk.CTkLabel(self.main_frame, text="密码", font=("Microsoft YaHei UI", 13, "bold"), text_color=self.theme['text_body'], anchor="w").grid(row=4, column=0, padx=30, sticky="ew")
        self.password_entry = ctk.CTkEntry(self.main_frame, font=("Microsoft YaHei UI", 14), show="*", placeholder_text="••••••••", corner_radius=8, height=40, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], text_color=self.theme['text_body'])
        self.password_entry.grid(row=5, column=0, padx=30, pady=(5, 10), sticky="ew")

        # --- Remember Password Checkbox ---
        self.remember_password_var = ctk.BooleanVar()
        self.remember_password_checkbox = ctk.CTkCheckBox(
            self.main_frame,
            text="记住密码",
            variable=self.remember_password_var,
            font=("Microsoft YaHei UI", 13),
            text_color=self.theme['text_light']
        )
        self.remember_password_checkbox.grid(row=6, column=0, padx=30, pady=(0, 15), sticky="w")
        
        self.status_label = ctk.CTkLabel(self.main_frame, text="", font=("Microsoft YaHei UI", 13), text_color=self.theme['danger'], wraplength=300)
        self.status_label.grid(row=7, column=0, padx=30, pady=(0, 10), sticky="ew")

        self.login_button = ctk.CTkButton(self.main_frame, text="登 录", command=self.handle_login, font=("Microsoft YaHei UI", 14, "bold"), height=45, corner_radius=8, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'])
        self.login_button.grid(row=8, column=0, padx=30, pady=(10, 10), sticky="ew")

        links_frame = ctk.CTkFrame(self.main_frame, fg_color="transparent")
        links_frame.grid(row=9, column=0, padx=30, pady=(0, 40))
        links_frame.grid_columnconfigure((0, 1), weight=1)

        self.register_button = ctk.CTkButton(links_frame, text="没有账户？立即注册", command=self.open_register_window, font=("Microsoft YaHei UI", 13), fg_color="transparent", text_color=self.theme['primary'], hover=False)
        self.register_button.grid(row=0, column=0, sticky="w")

        self.forgot_password_button = ctk.CTkButton(links_frame, text="忘记密码？", command=self.open_forgot_password_window, font=("Microsoft YaHei UI", 13), fg_color="transparent", text_color=self.theme['primary'], hover=False)
        self.forgot_password_button.grid(row=0, column=1, sticky="e")

        # --- Load Remembered Info ---
        self.email_entry.insert(0, self.parent.app_config.get("last_email", ""))
        if self.parent.app_config.get("remember_password"):
            self.remember_password_var.set(True)
            try:
                decoded_pwd = base64.b64decode(self.parent.app_config.get("last_password", "")).decode('utf-8')
                self.password_entry.insert(0, decoded_pwd)
            except Exception:
                # If decoding fails, just leave the password field blank
                self.parent.update_login_config("", "", False) # Clear corrupted password
                self.remember_password_var.set(False)


    def open_register_window(self):
        x = self.winfo_x(); y = self.winfo_y()
        self.withdraw()
        if self.register_win is None or not self.register_win.winfo_exists(): self.register_win = RegisterWindow(self, x, y)
        else: self.register_win.focus()
    def open_forgot_password_window(self):
        x = self.winfo_x(); y = self.winfo_y()
        self.withdraw()
        if self.forgot_password_win is None or not self.forgot_password_win.winfo_exists(): self.forgot_password_win = ForgotPasswordWindow(self, x, y)
        else: self.forgot_password_win.focus()
    def show_self(self, x, y): self.geometry(f"+{x}+{y}"); self.deiconify(); self.focus()
    def on_closing(self): self.parent.quit()
    def set_status(self, message, is_error=True):
        text_color = self.theme['danger'] if is_error else self.theme['primary']
        self.status_label.configure(text=message, text_color=text_color)
    def toggle_buttons(self, enabled=True):
        state = "normal" if enabled else "disabled"
        self.login_button.configure(state=state)
    def is_valid_email_format_client(self, email): return re.match(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$", email) is not None
    def handle_login(self):
        email = self.email_entry.get().lower(); password = self.password_entry.get()
        if not email or not password: self.set_status("邮箱和密码不能为空"); return
        if not self.is_valid_email_format_client(email): self.set_status("邮箱格式不正确"); return
        self.set_status("登录中...", is_error=False); self.toggle_buttons(False)
        threading.Thread(target=self._do_login, args=(email, password), daemon=True).start()
    def _do_login(self, email, password):
        try:
            response = requests.post(f"{SERVER_URL}/login", json={"email": email, "password": password}, timeout=10)
            data = response.json()
            if response.status_code == 200 and data.get("success") and "user" in data:
                self.parent.update_login_config(email, password, self.remember_password_var.get())
                self.user_data = data.get("user")
                self.after(0, self.set_status, "登录成功！", False)
                self.after(500, self.start_main_app)
            else:
                self.after(0, self.set_status, data.get("message", "登录失败")); self.after(0, self.toggle_buttons, True)
        except requests.RequestException: self.after(0, self.set_status, "无法连接到服务器"); self.after(0, self.toggle_buttons, True)
    def start_main_app(self):
        self.destroy()
        if self.register_win and self.register_win.winfo_exists(): self.register_win.destroy()
        if self.forgot_password_win and self.forgot_password_win.winfo_exists(): self.forgot_password_win.destroy()
        self.parent.on_login_success(self.user_data)

class PromptEditDialog(ctk.CTkToplevel):
    def __init__(self, parent, prompt_name="", prompt_content=""):
        super().__init__(parent)
        self.parent = parent; self.result = None; self.title("编辑提示词" if prompt_name else "添加新提示词")
        self.geometry("600x400"); self.transient(parent); self.grab_set()
        self.grid_columnconfigure(0, weight=1); self.grid_rowconfigure(2, weight=1)
        ctk.CTkLabel(self, text="提示词名称:").grid(row=0, column=0, padx=20, pady=(20, 5), sticky="w")
        self.name_entry = ctk.CTkEntry(self, font=("Microsoft YaHei UI", 13)); self.name_entry.grid(row=1, column=0, padx=20, pady=5, sticky="ew")
        self.name_entry.insert(0, prompt_name)
        ctk.CTkLabel(self, text="提示词内容 (使用 {context} 和 {requirements}):").grid(row=2, column=0, padx=20, pady=(10, 5), sticky="w")
        self.content_text = ctk.CTkTextbox(self, font=("Microsoft YaHei UI", 13), wrap="word"); self.content_text.grid(row=3, column=0, padx=20, pady=5, sticky="nsew")
        self.content_text.insert("1.0", prompt_content)
        button_frame = ctk.CTkFrame(self, fg_color="transparent"); button_frame.grid(row=4, column=0, padx=20, pady=20, sticky="e")
        self.save_button = ctk.CTkButton(button_frame, text="保 存", command=self._on_save); self.save_button.pack(side="right")
        self.cancel_button = ctk.CTkButton(button_frame, text="取 消", command=self.destroy, fg_color="transparent", border_width=1, text_color=parent.master.theme['text_body'], border_color=parent.master.theme['border']); self.cancel_button.pack(side="right", padx=10)
    def _on_save(self):
        name = self.name_entry.get().strip(); content = self.content_text.get("1.0", tk.END).strip()
        if not name or not content: messagebox.showwarning("输入错误", "不能为空", parent=self); return
        self.result = {"name": name, "content": content}; self.destroy()
    def get_input(self): self.wait_window(); return self.result

class PromptUploadDialog(ctk.CTkToplevel):
    def __init__(self, parent, prompt_name):
        super().__init__(parent)
        self.result = None; self.title("上传选项"); self.geometry("350x200"); self.transient(parent); self.grab_set()
        ctk.CTkLabel(self, text=f"请为 '{prompt_name}' 选择上传类型:", font=("Microsoft YaHei UI", 13)).pack(pady=20, padx=20)
        self.privacy_var = ctk.StringVar(value="private")
        ctk.CTkRadioButton(self, text="私密 (仅自己可见)", variable=self.privacy_var, value="private").pack(pady=5, padx=40, anchor="w")
        ctk.CTkRadioButton(self, text="公开 (其他用户可见)", variable=self.privacy_var, value="public").pack(pady=5, padx=40, anchor="w")
        button_frame = ctk.CTkFrame(self, fg_color="transparent"); button_frame.pack(pady=20, padx=20, fill="x"); button_frame.grid_columnconfigure((0,1), weight=1)
        self.ok_button = ctk.CTkButton(button_frame, text="确 定", command=self._on_ok); self.ok_button.grid(row=0, column=1, padx=5, sticky="ew")
        self.cancel_button = ctk.CTkButton(button_frame, text="取 消", command=self.destroy, fg_color="transparent", border_width=1); self.cancel_button.grid(row=0, column=0, padx=5, sticky="ew")
    def _on_ok(self): self.result = self.privacy_var.get(); self.destroy()
    def get_choice(self): self.wait_window(); return self.result

class PromptDownloadDialog(ctk.CTkToplevel):
    def __init__(self, parent, title, prompts_to_show: dict):
        super().__init__(parent); self.result = []; self.prompts_data = prompts_to_show; self.checkbox_vars = {}
        self.title(title); self.geometry("500x600"); self.transient(parent); self.grab_set()
        self.grid_rowconfigure(1, weight=1); self.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(self, text="请勾选您想要下载的提示词:", font=("Microsoft YaHei UI", 14, "bold")).grid(row=0, column=0, pady=15, padx=20, sticky="w")
        scroll_frame = ctk.CTkScrollableFrame(self); scroll_frame.grid(row=1, column=0, padx=20, sticky="nsew")
        if not self.prompts_data: ctk.CTkLabel(scroll_frame, text="没有可供下载的提示词。").pack(pady=20)
        else:
            for name, content in self.prompts_data.items():
                var = ctk.StringVar(value=""); cb = ctk.CTkCheckBox(scroll_frame, text=name, variable=var, onvalue=name, offvalue=""); cb.pack(anchor="w", padx=10, pady=8); self.checkbox_vars[name] = var
        button_frame = ctk.CTkFrame(self, fg_color="transparent"); button_frame.grid(row=2, column=0, padx=20, pady=20, sticky="ew"); button_frame.grid_columnconfigure((0,1), weight=1)
        self.ok_button = ctk.CTkButton(button_frame, text="下载选中项", command=self._on_ok, state="disabled" if not self.prompts_data else "normal"); self.ok_button.grid(row=0, column=1, padx=5, sticky="ew")
        self.cancel_button = ctk.CTkButton(button_frame, text="关 闭", command=self.destroy); self.cancel_button.grid(row=0, column=0, padx=5, sticky="ew")
    def _on_ok(self):
        self.result = [name for name, var in self.checkbox_vars.items() if var.get()]
        if not self.result: messagebox.showwarning("未选择", "您没有选择任何提示词。", parent=self); return
        self.destroy()
    def get_selection(self): self.wait_window(); return self.result

class PublicPromptsMarketDialog(ctk.CTkToplevel):
    def __init__(self, parent: Any, download_callback: Callable):
        super().__init__(parent); self.parent = parent; self.download_callback = download_callback; self.all_prompts = []; self.prompt_widgets = []
        self.title("公共提示词市场"); self.geometry("600x700"); self.transient(parent); self.grab_set()
        self.grid_columnconfigure(0, weight=1); self.grid_rowconfigure(1, weight=1)
        search_frame = ctk.CTkFrame(self, fg_color="transparent"); search_frame.grid(row=0, column=0, padx=20, pady=10, sticky="ew"); search_frame.grid_columnconfigure(0, weight=1)
        self.search_var = ctk.StringVar(); self.search_var.trace_add("write", self._filter_prompts)
        self.search_entry = ctk.CTkEntry(search_frame, textvariable=self.search_var, placeholder_text="搜索提示词名称或作者..."); self.search_entry.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        self.status_label = ctk.CTkLabel(search_frame, text="", width=20); self.status_label.grid(row=0, column=1)
        self.scroll_frame = ctk.CTkScrollableFrame(self); self.scroll_frame.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="nsew")
        threading.Thread(target=self._fetch_public_prompts, daemon=True).start()
    def _fetch_public_prompts(self):
        self.after(0, lambda: self.status_label.configure(text="⌛"))
        try:
            response = requests.get(f"{SERVER_URL}/cloud/prompts/public", timeout=15)
            if response.ok: self.all_prompts = response.json().get("prompts", []); self.after(0, self._populate_prompt_list)
            else: raise Exception(response.json().get("message", "未知错误"))
        except Exception as e: self.after(0, lambda: messagebox.showerror("获取失败", f"无法加载: {e}", parent=self))
        finally: self.after(0, lambda: self.status_label.configure(text=""))
    def _populate_prompt_list(self):
        for widget in self.prompt_widgets: widget.destroy()
        self.prompt_widgets.clear()
        if not self.all_prompts: ctk.CTkLabel(self.scroll_frame, text="目前没有可供下载的公共提示词。").pack(pady=30); return
        local_prompts = self.master.get_current_user_prompts()
        for prompt_data in self.all_prompts:
            card = ctk.CTkFrame(self.scroll_frame, border_width=1); card.pack(fill="x", padx=10, pady=5); card.grid_columnconfigure(0, weight=1)
            info_frame = ctk.CTkFrame(card, fg_color="transparent"); info_frame.grid(row=0, column=0, padx=15, pady=10, sticky="w")
            ctk.CTkLabel(info_frame, text=prompt_data.get("name", "未知"), font=("Microsoft YaHei UI", 14, "bold")).pack(side="left")
            ctk.CTkLabel(info_frame, text=f" (by {prompt_data.get('author', '匿名')})", font=("Microsoft YaHei UI", 12), text_color="gray").pack(side="left", padx=(5,0))
            is_downloaded = prompt_data.get("name") in local_prompts
            if is_downloaded: ctk.CTkButton(card, text="已下载", width=70, state="disabled").grid(row=0, column=1, padx=15, pady=10)
            else: ctk.CTkButton(card, text="下载", width=70, command=lambda data=prompt_data, c=card: self._download_prompt(data, c)).grid(row=0, column=1, padx=15, pady=10)
            self.prompt_widgets.append(card)
    def _filter_prompts(self, *args):
        query = self.search_var.get().lower()
        for i, prompt_data in enumerate(self.all_prompts):
            card = self.prompt_widgets[i]; name = prompt_data.get("name", "").lower(); author = prompt_data.get("author", "").lower()
            if query in name or query in author: card.pack(fill="x", padx=10, pady=5)
            else: card.pack_forget()
    def _download_prompt(self, prompt_data, card):
        self.download_callback(prompt_data, is_from_market=True)
        for widget in card.winfo_children():
            if isinstance(widget, ctk.CTkButton): widget.configure(text="已下载", state="disabled"); break

class CloudPromptManagerDialog(ctk.CTkToplevel):
    def __init__(self, parent: Any):
        super().__init__(parent); self.parent = parent; self.user_email = parent.current_user['email']
        self.title("管理我的云端提示词"); self.geometry("700x600"); self.transient(parent); self.grab_set()
        self.grid_columnconfigure(0, weight=1); self.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(self, text="管理云端提示词", font=("Microsoft YaHei UI", 14)).grid(row=0, column=0, padx=20, pady=(15, 5), sticky="w")
        self.scroll_frame = ctk.CTkScrollableFrame(self, label_text="我的云端提示词"); self.scroll_frame.grid(row=1, column=0, padx=20, pady=10, sticky="nsew")
        self.status_label = ctk.CTkLabel(self, text="正在加载..."); self.status_label.grid(row=2, column=0, padx=20, pady=10, sticky="w")
        ctk.CTkButton(self, text="关闭", command=self.destroy).grid(row=3, column=0, padx=20, pady=10, sticky="e")
        threading.Thread(target=self.load_prompts, daemon=True).start()
    def load_prompts(self):
        try:
            response = requests.get(f"{SERVER_URL}/cloud/prompts/user/{self.user_email}", timeout=15); response.raise_for_status(); data = response.json()
            if data.get("success"): self.after(0, self.populate_list, data.get("prompts", [])); self.after(0, lambda: self.status_label.configure(text=f"共 {len(data.get('prompts', []))} 个提示词"))
            else: raise Exception(data.get("message"))
        except Exception as e: self.after(0, lambda: self.status_label.configure(text=f"加载失败: {e}"))
    def populate_list(self, prompts):
        for widget in self.scroll_frame.winfo_children(): widget.destroy()
        if not prompts: ctk.CTkLabel(self.scroll_frame, text="云端暂无提示词").pack(pady=30); return
        for p in prompts:
            card = ctk.CTkFrame(self.scroll_frame); card.pack(fill="x", pady=5, padx=5); card.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(card, text=p.get("name", "未知"), font=("Microsoft YaHei UI", 13, "bold")).grid(row=0, column=0, padx=10, pady=5, sticky="w")
            bf = ctk.CTkFrame(card, fg_color="transparent"); bf.grid(row=0, column=1, padx=10, pady=5, sticky="e")
            ctk.CTkButton(bf, text="设为私密" if p.get("is_public") else "设为公开", width=80, command=lambda n=p.get("name"), s=p.get("is_public"): self.toggle_privacy(n, not s, card)).pack(side="left", padx=5)
            ctk.CTkButton(bf, text="删除", width=60, fg_color="#D9534F", hover_color="#C9302C", command=lambda n=p.get("name"): self.delete_prompt(n, card)).pack(side="left", padx=5)
    def toggle_privacy(self, name, new_is_public, card):
        threading.Thread(target=lambda: self._do_toggle(name, new_is_public), daemon=True).start()
    def _do_toggle(self, name, new_is_public):
        try: requests.patch(f"{SERVER_URL}/cloud/prompts/manage/{self.user_email}/{name}", json={"is_public": new_is_public}, timeout=10); self.after(100, self.load_prompts)
        except: pass
    def delete_prompt(self, name, card):
        if not messagebox.askyesno("确认", "确定删除吗？"): return
        card.pack_forget(); threading.Thread(target=lambda: requests.delete(f"{SERVER_URL}/cloud/prompts/manage/{self.user_email}/{name}", timeout=10), daemon=True).start()

def get_app_data_path():
    app_path = os.path.join(os.getenv('APPDATA') or os.path.expanduser("~"), "西门写作")
    os.makedirs(app_path, exist_ok=True); return app_path

def setup_logging(data_path):
    log_dir = os.path.join(data_path, "logs"); os.makedirs(log_dir, exist_ok=True)
    try:
        now = datetime.datetime.now(); seven_days_ago = now - datetime.timedelta(days=7)
        for f in os.listdir(log_dir):
            if f.startswith("app_") and f.endswith(".log"):
                try:
                    if datetime.datetime.strptime(f.replace("app_", "").replace(".log", ""), "%Y-%m-%d_%H-%M-%S") < seven_days_ago: os.remove(os.path.join(log_dir, f))
                except: pass
    except: pass
    log_file = os.path.join(log_dir, f"app_{datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.log")
    logging.basicConfig(filename=log_file, level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s', encoding='utf-8', force=True)
    return log_file

def load_app_config(data_path):
    config_path = os.path.join(data_path, "config.json")
    default_config = {
        "api_key": "", 
        "api_base_url": "https://api.openai.com/v1", 
        "ai_model": "gpt-4-turbo", 
        "appearance_mode": "System",
        "last_email": "",
        "remember_password": False,
        "last_password": ""
    }
    try:
        if os.path.exists(config_path):
            with open(config_path, 'r', encoding='utf-8') as f:
                c = json.load(f)
                for k, v in default_config.items(): c.setdefault(k, v)
                return c
    except: pass
    return default_config

def save_app_config(config, data_path):
    try:
        c = config.copy(); c.pop("prompts", None)
        with open(os.path.join(data_path, "config.json"), 'w', encoding='utf-8') as f: json.dump(c, f, ensure_ascii=False, indent=4)
    except: pass

def load_user_prompts(data_path, user_email):
    prompts_path = os.path.join(data_path, "prompts.json"); default = {"custom": {}, "market": {}}
    try:
        if os.path.exists(prompts_path):
            with open(prompts_path, 'r', encoding='utf-8') as f: all_p = json.load(f)
            u_p = all_p.get(user_email, {})
            u_p.setdefault("custom", {}); u_p.setdefault("market", {}); u_p.pop("system", None)
            return u_p
    except: pass
    return default

def save_user_prompts(data_path, user_email, prompts):
    prompts_path = os.path.join(data_path, "prompts.json"); all_p = {}
    try:
        if os.path.exists(prompts_path):
            with open(prompts_path, 'r', encoding='utf-8') as f: all_p = json.load(f)
    except: pass
    all_p[user_email] = prompts
    try:
        with open(prompts_path, 'w', encoding='utf-8') as f: json.dump(all_p, f, ensure_ascii=False, indent=4)
    except: pass

def generate_ai_content_stream(api_key, api_base, model, full_prompt, stop_flag_getter, stream_callback, final_callback):
    def api_call():
        try:
            if not api_key: raise ValueError("API Key 不能为空")
            client = openai.OpenAI(api_key=api_key, base_url=api_base)
            stream = client.chat.completions.create(model=model, messages=[{"role": "system", "content": "你是一位专业的作家..."}, {"role": "user", "content": full_prompt}], stream=True)
            full_res = []
            for chunk in stream:
                if stop_flag_getter(): break
                txt = chunk.choices[0].delta.content
                if txt: full_res.append(txt); stream_callback(txt)
            final_callback("".join(full_res), None)
        except Exception as e: final_callback(None, f"错误: {e}")
    threading.Thread(target=api_call, daemon=True).start()

def generate_ai_content_system_stream(email, model_id, full_prompt, stop_flag_getter, stream_callback, final_callback):
    def api_call():
        try:
            with requests.post(f"{SERVER_URL}/generate/system_stream", json={"email": email, "model_id": model_id, "prompt": full_prompt}, stream=True, timeout=120) as r:
                r.raise_for_status(); full_res = []
                for line in r.iter_lines():
                    if stop_flag_getter(): break
                    if line:
                        try:
                            d = json.loads(line.decode('utf-8'))
                            if 'error' in d: raise Exception(d['error'])
                            txt = d.get('chunk')
                            if txt: full_res.append(txt); stream_callback(txt)
                        except: pass
                if stop_flag_getter(): final_callback("".join(full_res), None, False)
                else: final_callback("".join(full_res), None, True)
        except Exception as e: final_callback(None, f"错误: {e}", False)
    threading.Thread(target=api_call, daemon=True).start()

def split_content_into_chapters(content):
    matches = list(re.finditer(r'^\s*(第\s*[一二三四五六七八九十百千万\d]+\s*章|Chapter\s*\d+)\s*.*$', content, re.MULTILINE))
    if not matches: return [("导入章节", content)]
    chapters = []
    for i, match in enumerate(matches):
        start = match.end(); end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
        chapters.append((match.group(0).strip(), content[start:end].strip()))
    return chapters

class AiWriterApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.app_data_path = get_app_data_path()
        self.log_file_path = setup_logging(self.app_data_path)
        self.app_config = load_app_config(self.app_data_path)
        
        ctk.set_appearance_mode(self.app_config.get("appearance_mode", "System"))
        ctk.set_default_color_theme("dark-blue")

        self.theme = {
            'bg': ("#F9FAFB", "#111827"),
            'sidebar': ("#FFFFFF", "#1F2937"),
            'card': ("#FFFFFF", "#1F2937"),
            'textbox_bg': ("#F3F4F6", "#374151"),
            'text_h1': ("#111827", "#F9FAFB"),
            'text_h2': ("#1F2937", "#E5E7EB"),
            'text_body': ("#374151", "#D1D5DB"),
            'text_light': ("#6B7280", "#9CA3AF"),
            'border': ("#E5E7EB", "#374151"),
            'primary': ("#4F46E5", "#7C3AED"),
            'primary_hover': ("#4338CA", "#6D28D9"),
            'primary_text': ("#FFFFFF", "#FFFFFF"),
            'danger': ("#EF4444", "#F87171"),
            'danger_hover': ("#DC2626", "#EF4444"),
            'secondary_btn': ("#E5E7EB", "#4B5563"),
            'secondary_btn_hover': ("#D1D5DB", "#374151"),
            'secondary_btn_text': ("#1F2937", "#F9FAFB"),
        }

        self.title("西门写作 Pro")
        self.geometry("1600x900")
        self.minsize(1600, 900)
        self.maxsize(1600, 900)
        self.configure(fg_color=self.theme['bg'])

        self.current_user = None; self.account_settings_window = None; self.cloud_books = set()
        self.selected_book_name = None; self.cloud_prompt_statuses = {}; self.user_prompts = {}; self.system_models = []
        self.books_path = None; self.settings_window = None; self.load_chapter_job = None
        self.book_buttons = {}; self.chapter_buttons = {}
        self.selected_chapter_filename = None
        self.context_checkbox_data = []
        self.stop_generation_flag = False

        self.use_smart_context_var = ctk.BooleanVar(value=True)
        self.use_manual_context_var = ctk.BooleanVar(value=False)
        self.recent_chars_count_var = ctk.StringVar(value="20000")
        self.system_model_var = ctk.StringVar()
        self.appearance_mode_var = ctk.StringVar(value=self.app_config.get("appearance_mode", "System"))

        self.FONT_TITLE = ctk.CTkFont(family="Microsoft YaHei UI", size=22, weight="bold")
        self.FONT_H2 = ctk.CTkFont(family="Microsoft YaHei UI", size=16, weight="bold")
        self.FONT_BODY = ctk.CTkFont(family="Microsoft YaHei UI", size=14)
        self.FONT_BUTTON = ctk.CTkFont(family="Microsoft YaHei UI", size=14, weight="bold")

        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)              
        self.grid_columnconfigure(2, weight=0, minsize=320) 
        self.grid_rowconfigure(0, weight=1)

        self._create_left_sidebar()
        self._create_center_main_area()
        self._create_right_sidebar()
        self.create_menu()

        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.bind("<Map>", self._on_map)
        self.after(100, self._update_context_controls_state)
        self.after(100, lambda: self.clear_chapter_view(clear_book_selection=True))
    
    def _on_map(self, event=None):
        """Handles the window becoming visible (e.g., de-minimizing)."""
        self.update()

    def update_login_config(self, email, password, remember_pwd):
        """Saves the user's login info based on their choices."""
        self.app_config["last_email"] = email
        self.app_config["remember_password"] = remember_pwd
        if remember_pwd:
            # WARNING: Storing passwords, even encoded, is not secure.
            # This is for convenience only in this specific application context.
            # In a real-world app, use a secure credential store like the system keyring.
            encoded_pwd = base64.b64encode(password.encode('utf-8')).decode('utf-8')
            self.app_config["last_password"] = encoded_pwd
        else:
            self.app_config["last_password"] = ""
        save_app_config(self.app_config, self.app_data_path)

    def get_current_user_prompts(self):
        c = self.user_prompts.get("custom", {}); m = self.user_prompts.get("market", {}); return {**c, **m}

    def on_login_success(self, user_data):
        self.current_user = user_data; self.username_label.configure(text=user_data.get('username', 'User'))
        self.system_models = user_data.get('system_models', [])
        user_hash = hashlib.md5(user_data['email'].encode()).hexdigest()
        self.books_path = os.path.join(self.app_data_path, "user_data", user_hash, "books")
        os.makedirs(self.books_path, exist_ok=True)
        self.user_prompts = load_user_prompts(self.app_data_path, user_data['email'])
        self._update_prompt_menu_options(); self._update_user_status_ui(); self._update_generation_controls()
        self.deiconify(); self.after(100, self.refresh_book_list)

    def on_closing(self):
        if self.current_user: save_user_prompts(self.app_data_path, self.current_user['email'], self.user_prompts)
        self.destroy(); self.quit()

    def _create_left_sidebar(self):
        self.left_sidebar_frame = ctk.CTkFrame(self, fg_color=self.theme['sidebar'], corner_radius=0)
        self.left_sidebar_frame.grid(row=0, column=0, sticky="nsew")
        self.left_sidebar_frame.grid_rowconfigure(2, weight=1)

        # --- Top Section ---
        top_frame = ctk.CTkFrame(self.left_sidebar_frame, fg_color="transparent")
        top_frame.grid(row=0, column=0, padx=20, pady=(30, 20), sticky="ew")
        ctk.CTkLabel(top_frame, text="📖 西门写作 Pro", font=self.FONT_TITLE, text_color=self.theme['text_h1']).pack(anchor="w")
        
        # --- User Info Card ---
        user_card = ctk.CTkFrame(self.left_sidebar_frame, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        user_card.grid(row=1, column=0, padx=15, pady=0, sticky="ew")
        user_card.grid_columnconfigure(0, weight=1)
        
        user_frame = ctk.CTkFrame(user_card, fg_color="transparent")
        user_frame.grid(row=0, column=0, padx=15, pady=10, sticky="ew")
        self.username_label = ctk.CTkLabel(user_frame, text="未登录", font=self.FONT_H2, text_color=self.theme['text_h2'])
        self.username_label.pack(side="left")
        ctk.CTkButton(user_frame, text="设置", width=50, height=24, font=ctk.CTkFont(size=12), fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'], command=self.open_account_settings_window).pack(side="right")
        
        status_frame = ctk.CTkFrame(user_card, fg_color="transparent")
        status_frame.grid(row=1, column=0, padx=15, pady=(0, 15), sticky="ew")
        self.membership_status_label = ctk.CTkLabel(status_frame, text="", font=ctk.CTkFont(size=12), text_color=self.theme['text_light']); self.membership_status_label.pack(anchor="w")
        self.word_balance_label = ctk.CTkLabel(status_frame, text="", font=ctk.CTkFont(size=12), text_color=self.theme['text_light']); self.word_balance_label.pack(anchor="w")

        # --- Book & Chapter Lists ---
        self.list_tabview = ctk.CTkTabview(self.left_sidebar_frame, fg_color="transparent", 
                                           segmented_button_selected_color=self.theme['primary'], 
                                           segmented_button_selected_hover_color=self.theme['primary_hover'],
                                           segmented_button_unselected_color=self.theme['sidebar'],
                                           text_color=self.theme['text_body'],
                                           segmented_button_fg_color=self.theme['card'],
                                           corner_radius=12)
        self.list_tabview.grid(row=2, column=0, padx=15, pady=15, sticky="nsew")
        self.list_tabview.add("书本"); self.list_tabview.add("章节")
        
        book_tab = self.list_tabview.tab("书本"); book_tab.grid_columnconfigure(0, weight=1); book_tab.grid_rowconfigure(0, weight=1)
        self.book_list_frame = ctk.CTkScrollableFrame(book_tab, fg_color="transparent"); self.book_list_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        
        book_ctrl = ctk.CTkFrame(book_tab, fg_color="transparent")
        book_ctrl.grid(row=1, column=0, sticky="ew", pady=5, padx=5)
        book_ctrl.grid_columnconfigure((0, 1), weight=1)
        ctk.CTkButton(book_ctrl, text="新建", height=32, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'], font=self.FONT_BODY, command=self.create_new_book).grid(row=0, column=0, padx=2, sticky="ew")
        ctk.CTkButton(book_ctrl, text="删除", height=32, fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], font=self.FONT_BODY, command=self.delete_selected_book).grid(row=0, column=1, padx=2, sticky="ew")

        chap_tab = self.list_tabview.tab("章节"); chap_tab.grid_columnconfigure(0, weight=1); chap_tab.grid_rowconfigure(0, weight=1)
        self.chapter_list_frame = ctk.CTkScrollableFrame(chap_tab, fg_color="transparent"); self.chapter_list_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        chap_ctrl = ctk.CTkFrame(chap_tab, fg_color="transparent"); chap_ctrl.grid(row=1, column=0, sticky="ew", pady=5, padx=5)
        chap_ctrl.grid_columnconfigure((0,1), weight=1)
        ctk.CTkButton(chap_ctrl, text="新建", height=32, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'], font=self.FONT_BODY, command=self.create_new_chapter).grid(row=0, column=0, padx=2, sticky="ew")
        ctk.CTkButton(chap_ctrl, text="删除", height=32, fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], font=self.FONT_BODY, command=self.delete_selected_chapter).grid(row=0, column=1, padx=2, sticky="ew")

        # --- Bottom Section ---
        bottom_frame = ctk.CTkFrame(self.left_sidebar_frame, fg_color="transparent")
        bottom_frame.grid(row=3, column=0, padx=20, pady=20, sticky="sew")
        
        ctk.CTkLabel(bottom_frame, text="外观模式:", font=ctk.CTkFont(size=12), text_color=self.theme['text_light']).pack(anchor="w")
        self.appearance_menu = ctk.CTkOptionMenu(bottom_frame, values=["Light", "Dark", "System"], variable=self.appearance_mode_var, command=self.change_appearance_mode, height=32, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], button_color=self.theme['secondary_btn_hover'], button_hover_color=self.theme['border'], corner_radius=8, font=self.FONT_BODY)
        self.appearance_menu.pack(fill="x", pady=(2, 10))

        ctk.CTkButton(bottom_frame, text="⚙️ 全局设置", height=36, command=self.open_settings_window, fg_color="transparent", border_width=1, border_color=self.theme['border'], text_color=self.theme['text_body'], font=self.FONT_BODY, corner_radius=8).pack(fill="x")

    def change_appearance_mode(self, new_mode):
        """ Triggers theme change and forces a refresh of dynamic UI elements. """
        ctk.set_appearance_mode(new_mode)
        self.app_config["appearance_mode"] = new_mode
        save_app_config(self.app_config, self.app_data_path)
        
        # Store selection before refresh
        selected_chap_before_refresh = self.selected_chapter_filename

        # Force refresh of lists which will destroy and recreate buttons
        self.refresh_book_list() 
        self.refresh_chapter_list()

        # Re-apply selection after refresh
        if self.selected_book_name and selected_chap_before_refresh:
            self.after(50, lambda: self.on_chapter_select_debounce(selected_chap_before_refresh))

    def _create_center_main_area(self):
        self.center_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.center_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 5), pady=10)
        self.center_frame.grid_rowconfigure(0, weight=1)
        self.center_frame.grid_columnconfigure(0, weight=1)

        self.main_editor_frame = ctk.CTkFrame(self.center_frame, corner_radius=12, fg_color=self.theme['card'], border_width=1, border_color=self.theme['border'])
        self.main_editor_frame.grid(row=0, column=0, sticky="nsew", pady=(0, 10))
        self.main_editor_frame.grid_rowconfigure(0, weight=1); self.main_editor_frame.grid_columnconfigure(0, weight=1)
        self.textbox = ctk.CTkTextbox(self.main_editor_frame, wrap="word", font=("Microsoft YaHei UI", 17), fg_color='transparent', text_color=self.theme['text_body'], corner_radius=8, border_width=0)
        self.textbox.grid(row=0, column=0, sticky="nsew", padx=20, pady=20)

        input_container = ctk.CTkFrame(self.center_frame, fg_color="transparent")
        input_container.grid(row=1, column=0, sticky="ew")
        input_container.grid_columnconfigure((0, 1), weight=1)

        ctx_frame = ctk.CTkFrame(input_container, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        ctx_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 5)); ctx_frame.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(ctx_frame, text="背景 {context}", font=self.FONT_H2, text_color=self.theme['text_h2']).pack(anchor="w", padx=15, pady=(10,5))
        self.background_textbox = ctk.CTkTextbox(ctx_frame, height=100, font=self.FONT_BODY, fg_color=self.theme['textbox_bg'], text_color=self.theme['text_body'], corner_radius=8, border_width=1, border_color=self.theme['border'])
        self.background_textbox.pack(fill="x", expand=True, padx=15, pady=(0, 15))

        req_frame = ctk.CTkFrame(input_container, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        req_frame.grid(row=0, column=1, sticky="nsew", padx=(5, 0)); req_frame.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(req_frame, text="要求 {requirements}", font=self.FONT_H2, text_color=self.theme['text_h2']).pack(anchor="w", padx=15, pady=(10,5))
        self.plot_textbox = ctk.CTkTextbox(req_frame, height=100, font=self.FONT_BODY, fg_color=self.theme['textbox_bg'], text_color=self.theme['text_body'], corner_radius=8, border_width=1, border_color=self.theme['border'])
        self.plot_textbox.pack(fill="x", expand=True, padx=15, pady=(0, 15))

        ctrl_bar = ctk.CTkFrame(self.center_frame, fg_color=self.theme['card'], height=70, corner_radius=12, border_width=1, border_color=self.theme['border'])
        ctrl_bar.grid(row=2, column=0, sticky="ew", pady=(10, 0))
        ctrl_bar.grid_columnconfigure(1, weight=1)

        # --- Left Group (Selectors) ---
        left_group = ctk.CTkFrame(ctrl_bar, fg_color="transparent")
        left_group.grid(row=0, column=0, padx=15, pady=10, sticky="w")

        ctk.CTkLabel(left_group, text="提示词:", text_color=self.theme['text_body'], font=self.FONT_BODY).grid(row=0, column=0, padx=(0, 8))
        self.prompt_menu_var = ctk.StringVar()
        self.prompt_menu = ctk.CTkOptionMenu(left_group, variable=self.prompt_menu_var, width=200, height=32, font=self.FONT_BODY, corner_radius=8, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], button_color=self.theme['secondary_btn_hover'], button_hover_color=self.theme['border'], dynamic_resizing=False)
        self.prompt_menu.grid(row=0, column=1)

        ctk.CTkLabel(left_group, text="模型:", text_color=self.theme['text_body'], font=self.FONT_BODY).grid(row=0, column=2, padx=(20, 8))
        self.model_selection_frame = ctk.CTkFrame(left_group, fg_color="transparent")
        self.model_selection_frame.grid(row=0, column=3)
        self.member_model_label = ctk.CTkLabel(self.model_selection_frame, text="", text_color=self.theme['primary'], font=self.FONT_BODY)
        self.system_model_menu = ctk.CTkOptionMenu(self.model_selection_frame, variable=self.system_model_var, width=160, height=32, font=self.FONT_BODY, corner_radius=8, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], button_color=self.theme['secondary_btn_hover'], button_hover_color=self.theme['border'], dynamic_resizing=False)
        
        # --- Right Group (Action Buttons) ---
        right_group = ctk.CTkFrame(ctrl_bar, fg_color="transparent")
        right_group.grid(row=0, column=2, padx=15, pady=10, sticky="e")
        
        self.save_chapter_button = ctk.CTkButton(right_group, text="保存", command=self.save_current_chapter, width=80, height=36, font=self.FONT_BODY, fg_color="transparent", border_width=1, border_color=self.theme['border'], text_color=self.theme['text_body'], hover_color=self.theme['secondary_btn_hover'])
        self.save_chapter_button.pack(side="left", padx=10)
        
        self.stop_generate_button = ctk.CTkButton(right_group, text="停止", command=self.stop_ai_generation, width=80, height=36, font=self.FONT_BODY, fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], state="disabled")
        self.stop_generate_button.pack(side="left", padx=0)
        
        self.generate_button = ctk.CTkButton(right_group, text="🚀 开始生成", command=self.start_ai_generation, width=120, height=40, font=self.FONT_BUTTON, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover'])
        self.generate_button.pack(side="left", padx=10)

    def _update_generation_controls(self):
        if not self.current_user: return

        membership_level = self.current_user.get('membership_level')
        has_api_access = membership_level in ['standard', 'premium']

        self.member_model_label.pack_forget()
        self.system_model_menu.pack_forget()
        
        models = []
        if has_api_access:
            models.append("自定义模型 (全局设置)")
        
        if self.system_models:
            for m in self.system_models:
                models.append(m.get('display_name', '未知模型'))
        else:
             if not has_api_access: models.append("无可用模型")

        self.system_model_menu.configure(values=models)
        if models: self.system_model_var.set(models[0])
        self.system_model_menu.pack(side="left")

    def _create_right_sidebar(self):
        self.right_sidebar_frame = ctk.CTkFrame(self, fg_color=self.theme['sidebar'], corner_radius=0, width=320)
        self.right_sidebar_frame.grid(row=0, column=2, sticky="nsew")
        self.right_sidebar_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(self.right_sidebar_frame, text="⚡ 助手", font=self.FONT_TITLE, text_color=self.theme['text_h1']).grid(row=0, column=0, padx=20, pady=(30, 15), sticky="w")

        ctx_card = ctk.CTkFrame(self.right_sidebar_frame, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        ctx_card.grid(row=1, column=0, sticky="ew", padx=15, pady=5); ctx_card.grid_columnconfigure(0, weight=1)
        
        self.smart_context_checkbox = ctk.CTkSwitch(ctx_card, text="智能关联上下文", variable=self.use_smart_context_var, font=self.FONT_H2, command=lambda: self._on_context_mode_switch('smart'), button_color=self.theme['primary'], progress_color=self.theme['primary_hover'])
        self.smart_context_checkbox.grid(row=0, column=0, padx=20, pady=20, sticky="w")
        
        self.smart_context_entry_frame = ctk.CTkFrame(ctx_card, fg_color="transparent")
        self.smart_context_entry_frame.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="w")
        ctk.CTkLabel(self.smart_context_entry_frame, text="关联字数:", text_color=self.theme['text_body'], font=self.FONT_BODY).pack(side="left")
        ctk.CTkEntry(self.smart_context_entry_frame, width=80, textvariable=self.recent_chars_count_var, fg_color=self.theme['textbox_bg'], border_color=self.theme['border'], font=self.FONT_BODY, corner_radius=8).pack(side="left", padx=10)

        man_card = ctk.CTkFrame(self.right_sidebar_frame, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        man_card.grid(row=2, column=0, sticky="ew", padx=15, pady=10); man_card.grid_columnconfigure(0, weight=1)
        
        self.manual_context_checkbox = ctk.CTkSwitch(man_card, text="手动选择章节", variable=self.use_manual_context_var, font=self.FONT_H2, command=lambda: self._on_context_mode_switch('manual'), button_color=self.theme['primary'], progress_color=self.theme['primary_hover'])
        self.manual_context_checkbox.grid(row=0, column=0, padx=20, pady=20, sticky="w")
        
        self.manual_context_frame = ctk.CTkFrame(man_card, fg_color="transparent")
        self.manual_context_frame.grid(row=1, column=0, padx=20, pady=(0, 20), sticky="ew")
        self.context_scroll_frame = ctk.CTkScrollableFrame(self.manual_context_frame, height=200, fg_color=self.theme['textbox_bg'], corner_radius=8, border_width=1, border_color=self.theme['border']); self.context_scroll_frame.pack(fill="x")
        
        btn_row = ctk.CTkFrame(self.manual_context_frame, fg_color="transparent"); btn_row.pack(fill="x", pady=(10,0))
        btn_row.grid_columnconfigure((0,1), weight=1)
        ctk.CTkButton(btn_row, text="全选", height=28, font=self.FONT_BODY, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'], command=self._toggle_all_context_checkboxes).grid(row=0, column=0, sticky='ew', padx=(0,5))
        ctk.CTkButton(btn_row, text="近5章", height=28, font=self.FONT_BODY, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'], command=self._select_last_five_chapters).grid(row=0, column=1, sticky='ew', padx=(5,0))

        tool_card = ctk.CTkFrame(self.right_sidebar_frame, fg_color=self.theme['card'], corner_radius=12, border_width=1, border_color=self.theme['border'])
        tool_card.grid(row=3, column=0, sticky="ew", padx=15, pady=5); tool_card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(tool_card, text="🛠️ 工具箱", font=self.FONT_H2, text_color=self.theme['text_h2']).pack(anchor="w", padx=20, pady=(20, 10))
        
        t_frame = ctk.CTkFrame(tool_card, fg_color="transparent"); t_frame.pack(fill="x", padx=15, pady=(0,15))
        ctk.CTkButton(t_frame, text="👤 角色起名", height=36, font=self.FONT_BODY, fg_color="transparent", border_width=1, border_color=self.theme['border'], text_color=self.theme['text_body'], hover_color=self.theme['secondary_btn_hover'], command=self._tool_generate_name).pack(fill="x", pady=4)
        ctk.CTkButton(t_frame, text="💡 灵感闪现", height=36, font=self.FONT_BODY, fg_color="transparent", border_width=1, border_color=self.theme['border'], text_color=self.theme['text_body'], hover_color=self.theme['secondary_btn_hover'], command=self._tool_get_plot_idea).pack(fill="x", pady=4)
        ctk.CTkButton(t_frame, text="📊 字数统计", height=36, font=self.FONT_BODY, fg_color="transparent", border_width=1, border_color=self.theme['border'], text_color=self.theme['text_body'], hover_color=self.theme['secondary_btn_hover'], command=self._tool_word_count).pack(fill="x", pady=4)

    def _update_user_status_ui(self):
        if not self.current_user: return
        
        level = self.current_user.get('membership_level')
        expiry_date = self.current_user.get('member_expiry_date', 'N/A')
        
        if level == 'premium':
            self.membership_status_label.configure(text=f"👑 高级会员 (到期: {expiry_date})", text_color="#FBBF24")
            self.word_balance_label.configure(text="∞ 无限字数")
        elif level == 'standard':
            self.membership_status_label.configure(text=f"⭐ 会员 (到期: {expiry_date})", text_color="#818CF8")
            self.word_balance_label.configure(text=f"剩余字数: {self.current_user.get('word_balance', 0)}")
        else:
            self.membership_status_label.configure(text="普通用户", text_color=self.theme['text_light'])
            self.word_balance_label.configure(text=f"剩余字数: {self.current_user.get('word_balance', 0)}")


    def _on_context_mode_switch(self, mode):
        if mode == 'smart' and self.use_smart_context_var.get():
            if self.use_manual_context_var.get(): self.use_manual_context_var.set(False); [v.set("") for v, c, fn in self.context_checkbox_data]
        elif mode == 'manual' and self.use_manual_context_var.get():
            if self.use_smart_context_var.get(): self.use_smart_context_var.set(False)
        self._update_context_controls_state()

    def _update_context_controls_state(self):
        is_smart = self.use_smart_context_var.get(); is_manual = self.use_manual_context_var.get()
        
        state = "normal" if is_smart else "disabled"
        for child in self.smart_context_entry_frame.winfo_children(): 
            child.configure(state=state)
            
        state = "normal" if is_manual else "disabled"
        for child in self.manual_context_frame.winfo_children(): 
            child.configure(state=state)

        if is_manual: self._update_context_checkbox_states()
        else: 
            for v, cb, fn in self.context_checkbox_data: cb.configure(state="disabled")

    def open_settings_window(self):
        if self.settings_window and self.settings_window.winfo_exists(): self.settings_window.focus(); return
        self.settings_window = ctk.CTkToplevel(self); win = self.settings_window
        win.title("全局设置"); win.geometry("700x800"); win.transient(self); win.grab_set()
        
        main_frame = ctk.CTkFrame(win, fg_color="transparent"); main_frame.pack(pady=20, padx=30, fill="both", expand=True)
        
        api_frame = ctk.CTkFrame(main_frame, fg_color=self.theme['card'], border_width=1, border_color=self.theme['border'], corner_radius=12)
        api_frame.pack(fill="x", pady=10)
        
        has_api_access = self.current_user.get('membership_level') in ['standard', 'premium']

        if has_api_access:
            api_frame.grid_columnconfigure(1, weight=1)
            ctk.CTkLabel(api_frame, text="API Key:", font=self.FONT_H2, text_color=self.theme['text_h2']).grid(row=0, column=0, padx=20, pady=15, sticky="w")
            win.api_key_entry = ctk.CTkEntry(api_frame, font=self.FONT_BODY, corner_radius=8); win.api_key_entry.grid(row=0, column=1, padx=20, pady=15, sticky="ew")
            win.api_key_entry.insert(0, self.app_config.get("api_key", ""))
            
            ctk.CTkLabel(api_frame, text="Base URL:", font=self.FONT_H2, text_color=self.theme['text_h2']).grid(row=1, column=0, padx=20, pady=15, sticky="w")
            win.api_url_entry = ctk.CTkEntry(api_frame, font=self.FONT_BODY, corner_radius=8); win.api_url_entry.grid(row=1, column=1, padx=20, pady=15, sticky="ew")
            win.api_url_entry.insert(0, self.app_config.get("api_base_url", ""))
            
            ctk.CTkLabel(api_frame, text="Model:", font=self.FONT_H2, text_color=self.theme['text_h2']).grid(row=2, column=0, padx=20, pady=15, sticky="w")
            win.model_entry = ctk.CTkEntry(api_frame, font=self.FONT_BODY, corner_radius=8); win.model_entry.grid(row=2, column=1, padx=20, pady=15, sticky="ew")
            win.model_entry.insert(0, self.app_config.get("ai_model", ""))
        else:
            ctk.CTkLabel(api_frame, text="会员可配置自定义 API。", text_color=self.theme['text_light'], font=self.FONT_BODY).pack(pady=30)
            win.api_key_entry = ctk.CTkEntry(win); win.api_url_entry = ctk.CTkEntry(win); win.model_entry = ctk.CTkEntry(win) 

        tab = ctk.CTkTabview(main_frame, corner_radius=12, border_width=1, border_color=self.theme['border'], segmented_button_selected_color=self.theme['primary'], segmented_button_unselected_color=self.theme['card']); tab.pack(fill="both", expand=True, pady=10)
        tab.add("本地提示词"); tab.add("市场提示词")
        win.custom_prompts_frame = ctk.CTkScrollableFrame(tab.tab("本地提示词"), fg_color="transparent"); win.custom_prompts_frame.pack(fill="both", expand=True, padx=10, pady=10)
        win.market_prompts_frame = ctk.CTkScrollableFrame(tab.tab("市场提示词"), fg_color="transparent"); win.market_prompts_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        ctrl = ctk.CTkFrame(main_frame, fg_color="transparent"); ctrl.pack(fill="x", pady=10)
        win.prompt_status_label = ctk.CTkLabel(ctrl, text=""); win.prompt_status_label.pack(side="left")
        win.add_prompt_btn = ctk.CTkButton(ctrl, text="新建", font=self.FONT_BODY, command=self._add_custom_prompt); win.add_prompt_btn.pack(side="right", padx=5)
        win.download_prompts_btn = ctk.CTkButton(ctrl, text="同步云端", font=self.FONT_BODY, command=self._handle_download_prompts); win.download_prompts_btn.pack(side="right", padx=5)
        win.more_prompts_btn = ctk.CTkButton(ctrl, text="市场", font=self.FONT_BODY, command=self._handle_more_prompts); win.more_prompts_btn.pack(side="right", padx=5)
        win.manage_cloud_prompts_btn = ctk.CTkButton(ctrl, text="管理云端", font=self.FONT_BODY, command=self._open_cloud_manager); win.manage_cloud_prompts_btn.pack(side="right", padx=5)

        ctk.CTkButton(main_frame, text="保存并关闭", command=self.save_settings, height=40, font=self.FONT_BUTTON, fg_color=self.theme['primary'], text_color=self.theme['primary_text'], hover_color=self.theme['primary_hover']).pack(fill="x", pady=10)
        threading.Thread(target=self._async_refresh_prompt_settings, daemon=True).start()

    def _set_prompt_controls_state(self, state):
        if self.settings_window and self.settings_window.winfo_exists():
            win = self.settings_window
            for btn in [win.add_prompt_btn, win.download_prompts_btn, win.more_prompts_btn, win.manage_cloud_prompts_btn]: btn.configure(state=state)
    def _async_refresh_prompt_settings(self):
        if not self.current_user or not self.settings_window: return
        self.after(0, lambda: self._set_prompt_controls_state("disabled"))
        try: self._api_get_cloud_prompt_statuses(); self.after(0, self._update_settings_prompt_lists_ui)
        finally: self.after(0, lambda: self._set_prompt_controls_state("normal"))
    def _update_settings_prompt_lists_ui(self):
        if not self.settings_window or not self.settings_window.winfo_exists(): return
        win = self.settings_window
        for w in win.custom_prompts_frame.winfo_children(): w.destroy()
        for name in sorted(self.user_prompts.get("custom", {}).keys()):
            f = ctk.CTkFrame(win.custom_prompts_frame, fg_color="transparent"); f.pack(fill="x", pady=2)
            status = self.cloud_prompt_statuses.get(name, "")
            st = " ☁️" if status else ""; ctk.CTkLabel(f, text=name+st, font=self.FONT_BODY).pack(side="left")
            ctk.CTkButton(f, text="删", width=30, fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], font=self.FONT_BODY, command=lambda n=name: self._delete_custom_prompt(n)).pack(side="right")
            ctk.CTkButton(f, text="编", width=30, fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], hover_color=self.theme['secondary_btn_hover'], font=self.FONT_BODY, command=lambda n=name: self._edit_custom_prompt(n)).pack(side="right", padx=5)
            if not status: ctk.CTkButton(f, text="⬆️", width=30, command=lambda n=name: self._handle_upload_single_prompt(n)).pack(side="right")
        for w in win.market_prompts_frame.winfo_children(): w.destroy()
        for name in sorted(self.user_prompts.get("market", {}).keys()):
            f = ctk.CTkFrame(win.market_prompts_frame, fg_color="transparent"); f.pack(fill="x", pady=2)
            ctk.CTkLabel(f, text=name, font=self.FONT_BODY).pack(side="left")
            ctk.CTkButton(f, text="删", width=30, fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], font=self.FONT_BODY, command=lambda n=name: self._delete_market_prompt(n)).pack(side="right")

    def _add_custom_prompt(self):
        res = PromptEditDialog(self.settings_window).get_input()
        if res:
             if res['name'] in self.get_current_user_prompts(): messagebox.showerror("错误", "名称已存在"); return
             self.user_prompts["custom"][res['name']] = res['content']; self._update_settings_prompt_lists_ui(); self._update_prompt_menu_options()
    def _edit_custom_prompt(self, name):
        res = PromptEditDialog(self.settings_window, name, self.user_prompts["custom"][name]).get_input()
        if res:
            if res['name'] != name and res['name'] in self.get_current_user_prompts(): messagebox.showerror("错误", "名称已存在"); return
            del self.user_prompts["custom"][name]; self.user_prompts["custom"][res['name']] = res['content']
            self._update_settings_prompt_lists_ui(); self._update_prompt_menu_options()
    def _delete_custom_prompt(self, name):
        if messagebox.askyesno("确认", f"删除 '{name}'?"):
            del self.user_prompts["custom"][name]; self._update_settings_prompt_lists_ui(); self._update_prompt_menu_options()
    def _delete_market_prompt(self, name):
        if messagebox.askyesno("确认", f"删除 '{name}'?"):
            del self.user_prompts["market"][name]; self._update_settings_prompt_lists_ui(); self._update_prompt_menu_options()

    def save_settings(self):
        if self.current_user.get('membership_level') in ['standard', 'premium']:
            self.app_config.update({'api_key': self.settings_window.api_key_entry.get(), 'api_base_url': self.settings_window.api_url_entry.get(), 'ai_model': self.settings_window.model_entry.get()})
        save_app_config(self.app_config, self.app_data_path); save_user_prompts(self.app_data_path, self.current_user['email'], self.user_prompts)
        self._update_generation_controls(); self._update_prompt_menu_options(); self.settings_window.destroy()

    def _update_prompt_menu_options(self):
        opts = [f"[本地] {k}" for k in sorted(self.user_prompts.get("custom", {}))] + [f"[市场] {k}" for k in sorted(self.user_prompts.get("market", {}))]
        if not opts: opts = ["默认"]
        self.prompt_menu.configure(values=opts); self.prompt_menu_var.set(opts[0])

    def refresh_book_list(self):
        if not self.books_path: return
        self._api_get_cloud_books()
        for w in self.book_list_frame.winfo_children(): w.destroy()
        self.book_buttons.clear()
        try:
            books = sorted([b for b in os.listdir(self.books_path) if os.path.isdir(os.path.join(self.books_path, b))])
        except FileNotFoundError:
            self.clear_chapter_view(clear_book_selection=True)
            return

        for b in books:
            if b == "__AIPrompts__": continue
            display_text = f"{b}  ☁️" if b in self.cloud_books else b
            
            btn = ctk.CTkButton(
                self.book_list_frame,
                text=display_text,
                font=self.FONT_BODY,
                anchor="w",
                command=lambda n=b: self.on_book_select(n),
                text_color=self.theme['text_body'],
                fg_color="transparent",
                hover_color=self.theme['secondary_btn_hover']
            )
            btn.pack(fill="x", pady=2, ipady=5)
            self.book_buttons[b] = btn

        self.update_book_selection_styles()
        
        if self.selected_book_name not in books:
            self.clear_chapter_view(clear_book_selection=True)

    def on_book_select(self, name):
        if self.selected_book_name == name:
            return
            
        self.selected_book_name = name
        self.update_book_selection_styles()
        self.refresh_chapter_list()

    def update_book_selection_styles(self):
        for name, button in self.book_buttons.items():
            if name == self.selected_book_name:
                button.configure(
                    fg_color=self.theme['primary'],
                    text_color=self.theme['primary_text'],
                    hover_color=self.theme['primary_hover']
                )
            else:
                button.configure(
                    fg_color="transparent",
                    text_color=self.theme['text_body'],
                    hover_color=self.theme['secondary_btn_hover']
                )

    def update_chapter_selection_styles(self):
        if not hasattr(self, 'chapter_buttons'): return
        for fn, btn in self.chapter_buttons.items():
            is_selected = fn == self.selected_chapter_filename
            fg_color = self.theme['primary'] if is_selected else "transparent"
            text_color = self.theme['primary_text'] if is_selected else self.theme['text_body']
            hover_color = self.theme['primary_hover'] if is_selected else self.theme['secondary_btn_hover']

            if fn == SETTINGS_FILENAME and not is_selected:
                text_color = self.theme['primary']

            btn.configure(fg_color=fg_color, text_color=text_color, hover_color=hover_color)

    def refresh_chapter_list(self):
        for w in self.chapter_list_frame.winfo_children(): w.destroy()
        self.chapter_buttons.clear()
        
        if self.load_chapter_job is None:
            self.selected_chapter_filename = None
            self.clear_chapter_view()
        
        path = os.path.join(self.books_path, self.selected_book_name) if self.selected_book_name else None
        if not (path and os.path.exists(path)):
            return

        settings_btn = ctk.CTkButton(self.chapter_list_frame, text="[ 本书设定 ] ✨", fg_color="transparent", text_color=self.theme['primary'], anchor="w", font=self.FONT_BODY, hover_color=self.theme['secondary_btn_hover'], command=lambda: self.on_chapter_select_debounce(SETTINGS_FILENAME))
        settings_btn.pack(fill="x", pady=1, ipady=5)
        self.chapter_buttons[SETTINGS_FILENAME] = settings_btn
        
        ctk.CTkFrame(self.chapter_list_frame, height=1, border_width=1, fg_color=self.theme['border']).pack(fill='x', padx=10, pady=8)

        chapters_files = sorted([f for f in os.listdir(path) if f.endswith(".txt") and f != SETTINGS_FILENAME])
        for filename in chapters_files:
            display_name = re.sub(r"^\d+_", "", filename).replace(".txt", "")
            try:
                with open(os.path.join(path, filename), 'r', encoding='utf-8') as f:
                    word_count = len(f.read())
            except Exception:
                word_count = 0
            
            display_text = f"{display_name}  ({word_count}字)"
            
            btn = ctk.CTkButton(self.chapter_list_frame, text=display_text, fg_color="transparent", text_color=self.theme['text_body'], anchor="w", font=self.FONT_BODY, hover_color=self.theme['secondary_btn_hover'], command=lambda fn=filename: self.on_chapter_select_debounce(fn))
            btn.pack(fill="x", pady=1, ipady=5)
            self.chapter_buttons[filename] = btn
        
        self.update_chapter_selection_styles()
        self._populate_context_panel()

    def _load_selected_chapter(self, filename):
        self.selected_chapter_filename = filename
        
        self.update_chapter_selection_styles()
        
        self.textbox.configure(state="normal", text_color=self.theme['text_body'])
        self.textbox.delete("1.0", tk.END)

        filepath = os.path.join(self.books_path, self.selected_book_name, filename)
        if filename == SETTINGS_FILENAME and not os.path.exists(filepath):
            try:
                with open(filepath, 'w', encoding='utf-8') as f:
                    f.write("# 在此输入本书的核心设定，例如世界观、主要角色、关键事件等。\n"
                            "# AI在生成任何章节时，都会自动读取并关联此处的全部内容。\n")
            except Exception:
                pass

        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                self.textbox.insert("1.0", f.read())
            self._update_context_checkbox_states()
        except Exception as e:
             self.textbox.insert("1.0", f"错误：无法加载文件。\n{e}")

    def on_chapter_select_debounce(self, filename):
        if self.load_chapter_job: self.after_cancel(self.load_chapter_job)
        self.load_chapter_job = self.after(50, lambda: self._load_selected_chapter(filename))
    
    def _populate_context_panel(self):
        for w in self.context_scroll_frame.winfo_children(): w.destroy()
        self.context_checkbox_data = []
        if not self.selected_book_name: return
        try:
            chapter_files = sorted([f for f in os.listdir(os.path.join(self.books_path, self.selected_book_name)) if f.endswith(".txt") and f != SETTINGS_FILENAME])
            for filename in chapter_files:
                display_name = re.sub(r"^\d+_", "", filename).replace(".txt", "")
                v = tk.StringVar(value="")
                cb = ctk.CTkCheckBox(self.context_scroll_frame, text=display_name, variable=v, onvalue=filename, offvalue="", font=self.FONT_BODY)
                cb.pack(anchor="w", padx=10, pady=5)
                self.context_checkbox_data.append((v, cb, filename))
        except: pass
        self._update_context_checkbox_states()
    
    def _update_context_checkbox_states(self):
        for v, cb, filename in self.context_checkbox_data:
            if filename == self.selected_chapter_filename:
                cb.configure(state="disabled"); v.set("")
            else:
                cb.configure(state="normal" if self.use_manual_context_var.get() else "disabled")
                
    def _toggle_all_context_checkboxes(self):
         if not self.use_manual_context_var.get(): self.use_manual_context_var.set(True); self._on_context_mode_switch('manual')
         is_any_selected = any(v.get() for v, c, fn in self.context_checkbox_data if c.cget("state")=="normal")
         for v, c, fn in self.context_checkbox_data:
             if c.cget("state")=="normal": v.set(c.cget("onvalue") if not is_any_selected else "")
             
    def _select_last_five_chapters(self):
        if not self.use_manual_context_var.get(): self.use_manual_context_var.set(True); self._on_context_mode_switch('manual')
        available_chapters = [data for data in self.context_checkbox_data if data[1].cget("state")=="normal"]
        for v, c, fn in self.context_checkbox_data: v.set("")
        for v, c, fn in available_chapters[-5:]: v.set(c.cget("onvalue"))

    def create_menu(self):
        m = tk.Menu(self); fm = tk.Menu(m, tearoff=0)
        fm.add_command(label="导入", command=self.import_book_from_file); fm.add_command(label="导出TXT", command=lambda: self.export_book('txt')); fm.add_separator(); fm.add_command(label="退出", command=self.on_closing)
        m.add_cascade(label="文件", menu=fm)
        cm = tk.Menu(m, tearoff=0); cm.add_command(label="上传当前书本", command=self.upload_current_book); cm.add_command(label="下载书本", command=self.show_download_dialog)
        m.add_cascade(label="云同步", menu=cm)
        self.configure(menu=m)

    def _api_get_cloud_books(self):
        try: self.cloud_books = set(requests.get(f"{SERVER_URL}/cloud/list_books/{self.current_user['email']}", timeout=5).json().get('cloud_books', []))
        except: pass
    def upload_current_book(self):
        if not self.selected_book_name: return
        if not messagebox.askyesno("确认", "上传会覆盖云端的版本，确定吗？"): return
        path = os.path.join(self.books_path, self.selected_book_name); chaps = []
        try:
            for f in os.listdir(path):
                if f.endswith(".txt"):
                    with open(os.path.join(path, f), 'r', encoding='utf-8') as fs: chaps.append({'filename': f, 'content': fs.read()})
            requests.post(f"{SERVER_URL}/cloud/upload_book/{self.current_user['email']}/{self.selected_book_name}", json={'chapters': chaps}, timeout=30)
            messagebox.showinfo("成功", "上传成功"); self.refresh_book_list()
        except Exception as e:
            messagebox.showerror("失败", f"上传失败: {e}")
    def show_download_dialog(self):
        self._api_get_cloud_books()
        try:
            local_books = {b for b in os.listdir(self.books_path) if os.path.isdir(os.path.join(self.books_path, b))}
        except FileNotFoundError:
            local_books = set()
        
        dl = sorted(list(self.cloud_books - local_books))
        if not dl: messagebox.showinfo("提示", "云端没有本地不存在的新书可供下载。"); return
        
        d = ctk.CTkToplevel(self); d.title("下载书本"); d.geometry("400x300"); d.transient(self); d.grab_set()
        ctk.CTkLabel(d, text="选择要下载的书:").pack(pady=10)
        sel = tk.StringVar()
        
        scroll_frame = ctk.CTkScrollableFrame(d); scroll_frame.pack(fill="both", expand=True, padx=20)
        for b in dl: ctk.CTkRadioButton(scroll_frame, text=b, variable=sel, value=b).pack(anchor="w", pady=2)
        
        ctk.CTkButton(d, text="下载选中项", command=lambda: [self.download_book(sel.get()), d.destroy()] if sel.get() else messagebox.showwarning("提示", "请选择一本书", parent=d)).pack(pady=10)

    def download_book(self, name):
        try:
            d = requests.get(f"{SERVER_URL}/cloud/download_book/{self.current_user['email']}/{name}", timeout=30).json()
            if d.get('success'):
                p = os.path.join(self.books_path, name); os.makedirs(p, exist_ok=True)
                for c in d['chapters']:
                     with open(os.path.join(p, c['filename']), 'w', encoding='utf-8') as f: f.write(c['content'])
                self.refresh_book_list()
            else:
                messagebox.showerror("下载失败", d.get("message", "未知错误"), parent=self)
        except Exception as e:
            messagebox.showerror("下载失败", f"网络错误: {e}", parent=self)
    
    def _api_get_cloud_prompt_statuses(self):
        try: self.cloud_prompt_statuses = requests.get(f"{SERVER_URL}/cloud/prompts/statuses/{self.current_user['email']}", timeout=5).json().get("statuses", {})
        except: pass
    def _handle_upload_single_prompt(self, name):
        pr = PromptUploadDialog(self.settings_window, name).get_choice()
        if pr: threading.Thread(target=lambda: self._api_upload_single_prompt(name, {"content": self.user_prompts["custom"][name], "public": pr=="public"}), daemon=True).start()
    def _api_upload_single_prompt(self, name, pl):
        try: requests.post(f"{SERVER_URL}/cloud/prompts/{self.current_user['email']}/{name}", json=pl, timeout=10); self._async_refresh_prompt_settings()
        except: pass
    def _handle_download_prompts(self):
        self._show_user_cloud_download_dialog()
    def _show_user_cloud_download_dialog(self):
        try:
            d = requests.get(f"{SERVER_URL}/cloud/prompts/user/{self.current_user['email']}", timeout=10).json().get("prompts", [])
            pd = {p["name"]: p["content"] for p in d if p["name"] not in self.get_current_user_prompts()}
            if pd: self._open_download_dialog("下载我的云端提示词", pd)
            else: messagebox.showinfo("提示", "没有可下载的新提示词。", parent=self.settings_window)
        except Exception as e:
            messagebox.showerror("错误", f"无法获取云端提示词: {e}", parent=self.settings_window)

    def _handle_more_prompts(self): PublicPromptsMarketDialog(self, lambda d, m: [self.user_prompts["market" if m else "custom"].update({d["name"]: d["content"]}), self._update_settings_prompt_lists_ui(), self._update_prompt_menu_options()])
    def _open_download_dialog(self, t, d):
        sel = PromptDownloadDialog(self.settings_window, t, d).get_selection()
        if sel: 
            for n in sel: self.user_prompts["custom"][n] = d[n]
            self._update_settings_prompt_lists_ui(); self._update_prompt_menu_options()
    def _open_cloud_manager(self): CloudPromptManagerDialog(self)

    def clear_chapter_view(self, clear_book_selection=False):
        if clear_book_selection:
            self.selected_book_name = None
        self.selected_chapter_filename = None
        
        self.textbox.configure(state="normal")
        self.textbox.delete("1.0", tk.END)

        if self.selected_book_name:
            self.textbox.insert("1.0", f"书本: 《{self.selected_book_name}》\n\n请从左侧“章节”列表中选择一项进行编辑。")
        else:
            self.textbox.insert("1.0", "欢迎使用西门写作！\n\n请从左侧“书本”列表选择或新建一本书籍开始您的创作。")
        
        self.textbox.configure(state="disabled", text_color=self.theme['text_light'])

        if self.chapter_list_frame:
            for w in self.chapter_list_frame.winfo_children(): w.destroy()
        
        if clear_book_selection:
            self._populate_context_panel()
            self.update_book_selection_styles()
            if self.books_path is None: self.refresh_book_list()

    def create_new_book(self):
        if not self.current_user: messagebox.showwarning("提示", "请先登录"); return
        n = simpledialog.askstring("新建书本", "请输入书名:", parent=self)
        if n and n.strip():
            n = n.strip()
            path = os.path.join(self.books_path, n)
            if os.path.exists(path):
                messagebox.showerror("错误", f"名为 '{n}' 的书本已存在。", parent=self)
                return
            os.makedirs(path, exist_ok=True)
            self.refresh_book_list()
            self.on_book_select(n)
        elif n is not None:
             messagebox.showwarning("输入错误", "书名不能为空。", parent=self)

    def create_new_chapter(self):
        if not self.selected_book_name:
            messagebox.showwarning("提示", "请先选择一本书。", parent=self)
            return
        p = os.path.join(self.books_path, self.selected_book_name)
        n = len([f for f in os.listdir(p) if f.endswith(".txt") and f != SETTINGS_FILENAME]) + 1
        chap_name = f"{n:03d}_第{n}章.txt"
        chap_title = f"# 第{n}章\n\n"
        try:
            with open(os.path.join(p, chap_name), 'w', encoding='utf-8') as f: f.write(chap_title)
            self.refresh_chapter_list()
            self.after(100, lambda: self.on_chapter_select_debounce(chap_name))
        except Exception as e:
            messagebox.showerror("错误", f"创建章节失败: {e}", parent=self)

    def delete_selected_book(self):
        if self.selected_book_name and messagebox.askyesno("删除确认", f"您确定要永久删除书本 '{self.selected_book_name}' 及其所有章节和设定吗？\n此操作无法撤销！", parent=self):
            try:
                shutil.rmtree(os.path.join(self.books_path, self.selected_book_name))
                self.clear_chapter_view(clear_book_selection=True)
                self.refresh_book_list()
            except Exception as e:
                messagebox.showerror("删除失败", f"无法删除书本: {e}", parent=self)

    def delete_selected_chapter(self):
        if not self.selected_chapter_filename or not self.selected_book_name:
            messagebox.showwarning("提示", "请先选择一个章节。", parent=self); return
            
        if self.selected_chapter_filename == SETTINGS_FILENAME:
            messagebox.showerror("操作无效", "无法删除 [本书设定]，此为受保护文件。", parent=self)
            return
            
        if messagebox.askyesno("删除确认", f"您确定要删除这个章节吗？", parent=self):
            try:
                os.remove(os.path.join(self.books_path, self.selected_book_name, self.selected_chapter_filename))
                self.refresh_chapter_list()
            except Exception as e:
                messagebox.showerror("删除失败", f"无法删除章节: {e}", parent=self)

    def save_current_chapter(self):
        if self.selected_chapter_filename and self.selected_book_name:
            self.save_chapter_button.configure(text="保存中...")
            try:
                filepath = os.path.join(self.books_path, self.selected_book_name, self.selected_chapter_filename)
                with open(filepath, 'w', encoding='utf-8') as f:
                    f.write(self.textbox.get("1.0", tk.END))
                
                self.save_chapter_button.configure(text="✅ 已保存", text_color=("#22C55E", "#4ADE80"))
                self.after(2000, lambda: self.save_chapter_button.configure(text="保存", text_color=self.theme['text_body']))
                
                self.after(100, self.refresh_chapter_list)
                self.after(200, lambda: self.on_chapter_select_debounce(self.selected_chapter_filename))

            except Exception as e:
                messagebox.showerror("保存失败", f"无法保存文件: {e}", parent=self)
                self.save_chapter_button.configure(text="保存")
        else:
             messagebox.showwarning("提示", "没有选中的章节或设定可以保存。", parent=self)
    
    def import_book_from_file(self):
        if not self.books_path: return
        fp = filedialog.askopenfilename(filetypes=(("Text/Word", "*.txt;*.docx"),), parent=self)
        if not fp: return
        bn = os.path.splitext(os.path.basename(fp))[0]; p = os.path.join(self.books_path, bn)
        if os.path.exists(p): messagebox.showerror("错误", "同名书本已存在", parent=self); return
        try:
            c = docx.Document(fp).paragraphs if fp.endswith('.docx') else open(fp, 'r', encoding='utf-8').read()
            c = "\n".join([x.text for x in c]) if isinstance(c, list) else c
            os.makedirs(p); chaps = split_content_into_chapters(c)
            for i, (t, txt) in enumerate(chaps):
                safe_t = re.sub(r'[\\/*?:"<>|]', '_', t[:20])
                with open(os.path.join(p, f"{i+1:03d}_{safe_t}.txt"), 'w', encoding='utf-8') as f: f.write(f"# {t}\n\n{txt}")
            self.refresh_book_list(); messagebox.showinfo("成功", "导入成功", parent=self)
        except Exception as e: messagebox.showerror("失败", str(e), parent=self)
    
    def export_book(self, fmt):
        if not self.selected_book_name: messagebox.showwarning("提示", "请先选择要导出的书本。", parent=self); return
        fp = filedialog.asksaveasfilename(defaultextension=f".{fmt}", initialfile=self.selected_book_name, filetypes=(("Text file", "*.txt"), ("All files", "*.*")), parent=self)
        if not fp: return
        
        try:
            content = []
            book_path = os.path.join(self.books_path, self.selected_book_name)
            chapter_files = sorted([f for f in os.listdir(book_path) if f.endswith(".txt") and f != SETTINGS_FILENAME])
            for chap_file in chapter_files:
                with open(os.path.join(book_path, chap_file), 'r', encoding='utf-8') as f:
                    content.append(f.read())
            
            with open(fp, 'w', encoding='utf-8') as f:
                f.write("\n\n\n".join(content))
            messagebox.showinfo("成功", f"书本已成功导出到 {fp}", parent=self)
        except Exception as e:
            messagebox.showerror("导出失败", f"发生错误: {e}", parent=self)

    def stop_ai_generation(self): self.stop_generation_flag = True
    
    def start_ai_generation(self):
        if self.generate_button.cget("state") == "disabled": return
        
        if not self.selected_chapter_filename or not self.selected_book_name:
            messagebox.showwarning("操作提示", "请先从左侧列表中选择一本书和一个章节（或本书设定），然后再开始生成。", parent=self)
            return

        self.stop_generation_flag = False
        self.generate_button.configure(state="disabled", text="生成中...")
        self.stop_generate_button.configure(state="normal")
        
        book_settings = ""
        settings_path = os.path.join(self.books_path, self.selected_book_name, SETTINGS_FILENAME)
        try:
            if os.path.exists(settings_path):
                with open(settings_path, 'r', encoding='utf-8') as f:
                    settings_content = f.read().strip()
                    if settings_content:
                        settings_content = "\n".join([line for line in settings_content.splitlines() if not line.strip().startswith('#')])
                        if settings_content.strip():
                           book_settings = f"[本书核心设定]:\n{settings_content}\n\n---\n\n"
        except Exception:
            pass 

        ctx = book_settings
        
        background_ctx = self.background_textbox.get("1.0", "end-1c").strip()
        if background_ctx:
             ctx += f"[补充背景]:\n{background_ctx}\n\n---\n\n"
        else:
            p = os.path.join(self.books_path, self.selected_book_name); cp = []
            if self.use_smart_context_var.get():
                try:
                     lim = int(self.recent_chars_count_var.get())
                     all_chaps = sorted([fn for fn in self.chapter_buttons.keys() if fn != SETTINGS_FILENAME])
                     if self.selected_chapter_filename in all_chaps:
                         chapters_before = all_chaps[:all_chaps.index(self.selected_chapter_filename)]
                         for c in reversed(chapters_before):
                             if len("".join(cp)) > lim: break
                             with open(os.path.join(p, c), 'r', encoding='utf-8') as f: cp.insert(0, f.read() + "\n---\n")
                except Exception: pass
            elif self.use_manual_context_var.get():
                for v, cb, filename in self.context_checkbox_data:
                    if v.get(): 
                        try:
                            with open(os.path.join(p, filename), 'r', encoding='utf-8') as f: cp.append(f.read())
                        except Exception: pass
            if cp:
                ctx += "[关联章节内容]:\n" + "".join(cp)

        ctx += "\n\n[当前章节正文]:\n" + self.textbox.get("1.0", tk.END)
        req = self.plot_textbox.get("1.0", "end-1c").strip()
        
        p_sel = self.prompt_menu_var.get()
        tpl = "请根据我提供的背景(Context)和要求(Requirements)，续写下面的故事。要求自然、流畅、符合人设和逻辑。\n\nContext:\n{context}\n\nRequirements:\n{requirements}"
        if "[本地]" in p_sel: tpl = self.user_prompts["custom"].get(p_sel.split("] ")[1], tpl)
        elif "[市场]" in p_sel: tpl = self.user_prompts["market"].get(p_sel.split("] ")[1], tpl)
        prompt = tpl.format(context=ctx, requirements=req)

        sel_model = self.system_model_var.get()
        membership_level = self.current_user.get('membership_level')
        use_custom_model = membership_level in ['standard', 'premium'] and "自定义" in sel_model

        if use_custom_model:
            generate_ai_content_stream(self.app_config["api_key"], self.app_config["api_base_url"], self.app_config["ai_model"], prompt, lambda: self.stop_generation_flag, self._stream, self._on_gen_end)
        else:
            mid = next((m['id'] for m in self.system_models if m['display_name'] == sel_model), None)
            if not mid: 
                if self.system_models:
                    mid = self.system_models[0]['id']
                    messagebox.showwarning("模型提示", f"选择的模型'{sel_model}'无效，已自动为您切换到'{self.system_models[0]['display_name']}'。", parent=self)
                else:
                    messagebox.showerror("错误", "没有可用的系统模型。请联系管理员或成为会员使用自定义模型。", parent=self)
                    self._on_gen_end(None, "无可用模型", False)
                    return
            generate_ai_content_system_stream(self.current_user['email'], mid, prompt, lambda: self.stop_generation_flag, self._stream, self._on_gen_end)

    def _stream(self, c):
        self.textbox.insert(tk.END, c)
        self.textbox.see(tk.END)

    def _on_gen_end(self, res, err, sys=False):
        self.generate_button.configure(state="normal", text="🚀 开始生成")
        self.stop_generate_button.configure(state="disabled")
        if err:
            self.textbox.insert(tk.END, f"\n\n[发生错误: {err}]")
        elif self.stop_generation_flag:
            self.textbox.insert(tk.END, "\n\n[生成被用户手动停止]")
        
        if not err and not self.stop_generation_flag:
            # Word deduction logic: Deduct if it was a system call AND the user is NOT a premium member
            if sys and res and self.current_user.get('membership_level') != 'premium':
                threading.Thread(target=lambda: requests.post(f"{SERVER_URL}/user/deduct_words", json={"email": self.current_user['email'], "count": len(res)}), daemon=True).start()
                # We might want to fetch the user profile again to update the word count UI
                self.after(1000, self._fetch_profile) 

            self.save_current_chapter()

    def open_account_settings_window(self):
        if not self.current_user: return
        if self.account_settings_window and self.account_settings_window.winfo_exists(): self.account_settings_window.focus(); return
        self.account_settings_window = ctk.CTkToplevel(self); w = self.account_settings_window
        w.title("账户设置"); w.geometry("400x500"); w.transient(self); w.grab_set()
        f = ctk.CTkFrame(w, fg_color="transparent"); f.pack(pady=20, padx=20, fill="both", expand=True)
        
        ctk.CTkLabel(f, text="用户名").pack(anchor="w"); u = ctk.CTkEntry(f); u.pack(fill="x", pady=(0,10)); u.insert(0, self.current_user.get('username',''))
        ctk.CTkButton(f, text="更新用户名", command=lambda: self.save_account_settings(u.get(), None, None, None)).pack(fill="x")
        
        ctk.CTkLabel(f, text="更改密码").pack(anchor="w", pady=(20,0))
        op = ctk.CTkEntry(f, placeholder_text="当前密码", show="*"); op.pack(fill="x", pady=5)
        np = ctk.CTkEntry(f, placeholder_text="新密码", show="*"); np.pack(fill="x", pady=5)
        ctk.CTkButton(f, text="更新密码", command=lambda: self.save_account_settings(None, op.get(), np.get(), np.get())).pack(fill="x", pady=5)
        
        ctk.CTkButton(f, text="兑换激活码", fg_color=self.theme['secondary_btn'], text_color=self.theme['secondary_btn_text'], command=self._redeem_key_dialog).pack(fill="x", pady=20)
        ctk.CTkButton(f, text="退出登录", fg_color=self.theme['danger'], hover_color=self.theme['danger_hover'], command=self.handle_logout).pack(fill="x")

    def _redeem_key_dialog(self):
        k = simpledialog.askstring("兑换中心", "请输入您的激活码:", parent=self.account_settings_window)
        if k: threading.Thread(target=self._do_redeem, args=(k,), daemon=True).start()

    def _do_redeem(self, k):
        try:
            r = requests.post(f"{SERVER_URL}/user/redeem", json={"email": self.current_user['email'], "key": k}).json()
            if r.get("success"):
                messagebox.showinfo("成功", "兑换成功！您的账户信息已更新。", parent=self.account_settings_window)
                self._fetch_profile()
            else: messagebox.showerror("失败", r.get("message", "未知错误"), parent=self.account_settings_window)
        except Exception as e:
            messagebox.showerror("网络错误", f"无法连接到服务器: {e}", parent=self.account_settings_window)

    def _fetch_profile(self):
        try:
            r = requests.get(f"{SERVER_URL}/user/profile/{self.current_user['email']}").json()
            if r.get("success"):
                self.current_user = r.get("user")
                self.after(0, self._update_user_status_ui)
                self.after(0, self._update_generation_controls)
        except: pass

    def save_account_settings(self, nu, op, np, cp):
        if nu: requests.post(f"{SERVER_URL}/update_profile", json={'email': self.current_user['email'], 'new_username': nu})
        if op and np: requests.post(f"{SERVER_URL}/change_password", json={'email': self.current_user['email'], 'old_password': op, 'new_password': np})
        self.account_settings_window.destroy(); messagebox.showinfo("提示", "账户更新请求已发送", parent=self)

    def handle_logout(self):
        if messagebox.askyesno("退出登录", "您确定要退出当前账户吗？"):
            if self.current_user: save_user_prompts(self.app_data_path, self.current_user['email'], self.user_prompts)
            self.destroy(); global restart_app; restart_app = True

    def _tool_word_count(self): 
        if self.selected_chapter_filename:
            word_count = len(self.textbox.get('1.0', tk.END))
            messagebox.showinfo("字数统计", f"当前章节总字数: {word_count}", parent=self)
        else:
            messagebox.showinfo("字数统计", "请先选择一个章节。", parent=self)

    def _tool_get_plot_idea(self): messagebox.showinfo("灵感闪现", random.choice(["一个突如其来的访客，带来了一个埋藏多年的秘密。", "主角意外发现了一件能够穿越到过去的物品。", "在一场百年不遇的暴风雨中，所有的通讯都中断了，而岛上的居民开始一个个消失。", "一个看似普通的古董，竟然是某个失落文明的关键。", "主角收到了一个已故亲人的来信，信中揭示了一个惊人的家族真相。"]), parent=self)
    def _tool_generate_name(self):
        ln_cn = ["李", "王", "张", "刘", "陈", "杨", "赵", "黄", "周", "吴"]
        fn_cn_m = ["伟", "磊", "勇", "杰", "强", "军", "波", "浩", "明", "龙"]
        fn_cn_f = ["芳", "娜", "敏", "静", "丽", "艳", "燕", "娟", "萍", "婷"]
        ln_en = ["Smith", "Jones", "Williams", "Brown", "Davis", "Miller", "Wilson", "Moore"]
        fn_en_m = ["James", "John", "Robert", "Michael", "William", "David", "Richard", "Joseph"]
        fn_en_f = ["Mary", "Patricia", "Jennifer", "Linda", "Elizabeth", "Barbara", "Susan", "Jessica"]
        
        name_cn_m = f"{random.choice(ln_cn)}{random.choice(fn_cn_m)}"
        name_cn_f = f"{random.choice(ln_cn)}{random.choice(fn_cn_f)}"
        name_en_m = f"{random.choice(fn_en_m)} {random.choice(ln_en)}"
        name_en_f = f"{random.choice(fn_en_f)} {random.choice(ln_en)}"

        messagebox.showinfo("随机起名", f"中文男: {name_cn_m}\n中文女: {name_cn_f}\n\n英文男: {name_en_m}\n英文女: {name_en_f}", parent=self)

def start_app():
    root = AiWriterApp()
    root.withdraw()
    login = LoginWindow(root)
    root.mainloop()

if __name__ == "__main__":
    restart_app = True
    while restart_app:
        restart_app = False
        start_app()